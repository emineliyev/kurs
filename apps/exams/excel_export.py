"""
Results workbook (.xlsx) of one exam — or of the same exam given to several classes, in one file:

  1. Nəticələr      — a summary on top, then every student ranked by score
  2. Cavablar       — the answer key row, then each student's answer to every question, green / red
  3. Sual təhlili   — per question: how many got it right, how the choices were spread, how hard it was

Questions and choice letters follow the test's own order (A = the first choice as the teacher wrote
it), so students who saw a shuffled exam still line up with each other and with the key.
"""
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.tests_app.models import Question
from .models import ExamResult

LETTERS = 'ABCDEFGH'

# ── look ──────────────────────────────────────────────────────────────────────
NAVY = '1E3A8A'
INK = '0F172A'
MUTED = '64748B'
LINE = 'CBD5E1'
ZEBRA = 'F8FAFC'
GREEN, GREEN_BG, GREEN_DARK = '15803D', 'DCFCE7', '166534'
RED, RED_BG = 'DC2626', 'FEE2E2'
AMBER, AMBER_BG = 'B45309', 'FEF3C7'
BLUE_BG = 'DBEAFE'

FONT = 'Calibri'
thin = Side(style='thin', color=LINE)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal='center', vertical='center')
LEFT = Alignment(horizontal='left', vertical='center')
WRAP_CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)


def _fill(color):
    return PatternFill('solid', start_color=color, end_color=color)


def _font(color=INK, bold=False, size=11, italic=False):
    return Font(name=FONT, color=color, bold=bold, size=size, italic=italic)


def _header(ws, row, values, start_col=1):
    for i, value in enumerate(values):
        c = ws.cell(row=row, column=start_col + i, value=value)
        c.fill = _fill(NAVY)
        c.font = _font('FFFFFF', bold=True)
        c.alignment = WRAP_CENTER
        c.border = BOX


def _title_block(ws, width, title, lines):
    """Navy title bar over `width` columns, then grey info lines under it."""
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    c = ws.cell(row=1, column=1, value=title)
    c.fill = _fill(NAVY)
    c.font = _font('FFFFFF', bold=True, size=15)
    c.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws.row_dimensions[1].height = 30
    for i, text in enumerate(lines, start=2):
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=width)
        c = ws.cell(row=i, column=1, value=text)
        c.font = _font(MUTED, size=10, bold=(i == 2))
        c.alignment = Alignment(horizontal='left', vertical='center', indent=1)
        c.fill = _fill(ZEBRA)
    return len(lines) + 3   # first free row (one blank row after the block)


def _widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _num(value):
    if value is None:
        return None
    value = round(float(value), 2)
    return int(value) if value == int(value) else value


# ── data ──────────────────────────────────────────────────────────────────────
def _collect(exams):
    test = exams[0].test
    questions = [q for q in test.questions.prefetch_related('answers').order_by('order', 'id')
                 if q.question_type != Question.TYPE_PASSAGE]
    choices = {q.id: list(q.answers.all()) for q in questions}
    # submitted results, plus anyone still writing (listed at the bottom, without answers)
    results = (ExamResult.objects.filter(exam__in=exams)
               .select_related('student', 'exam__class_room')
               .prefetch_related('answers'))
    rows = []
    for r in results:
        given = {a.question_id: a for a in r.answers.all()}
        cells, correct, wrong, empty, pending = [], 0, 0, 0, 0
        for q in questions:
            a = given.get(q.id)
            text, state = _answer_text(q, choices[q.id], a)
            cells.append((text, state))
            if state == 'right':
                correct += 1
            elif state in ('wrong', 'part'):
                wrong += 1
            elif state == 'pending':
                pending += 1
            else:
                empty += 1
        rows.append({
            'result': r,
            'name': r.student.get_full_name() or r.student.username,
            'username': r.student.username,
            'class': r.exam.class_room.name,
            'correct': correct, 'wrong': wrong, 'empty': empty, 'pending': pending,
            'cells': cells, 'given': given,
        })
    done = [x for x in rows if x['result'].status == ExamResult.STATUS_SUBMITTED]
    done.sort(key=lambda x: (-(x['result'].score or 0), x['name'].lower()))
    rest = sorted((x for x in rows if x not in done), key=lambda x: x['name'].lower())
    return test, questions, choices, done, rest


def _answer_text(q, q_choices, a):
    """(what to show, state) — state: right / wrong / part / pending / blank."""
    if a is None:
        return '', 'blank'
    qtype = q.question_type
    if qtype in Question.MCQ_TYPES:
        if not a.selected_answer_id:
            return '', 'blank'
        letter = next((LETTERS[i] for i, c in enumerate(q_choices) if c.id == a.selected_answer_id and i < len(LETTERS)), '?')
        return letter, 'right' if a.is_correct else 'wrong'
    text = (a.answer_text or a.open_text or '').strip()
    if qtype in Question.MANUAL_GRADE_TYPES:
        if not text:
            return '', 'blank'
        if a.is_correct is None:
            return '?', 'pending'
        if a.is_correct:
            return '✓', 'right'
        return ('½', 'part') if a.points_earned > 0 else ('✗', 'wrong')
    if not text:
        return '', 'blank'
    return (text if len(text) <= 10 else text[:9] + '…'), 'right' if a.is_correct else 'wrong'


def _key_text(q, q_choices):
    if q.question_type in Question.MCQ_TYPES:
        return ''.join(LETTERS[i] for i, c in enumerate(q_choices) if c.is_correct and i < len(LETTERS))
    if q.question_type in Question.MANUAL_GRADE_TYPES:
        return '✎'
    key = (q.correct_answer or '').strip()
    return key if len(key) <= 10 else key[:9] + '…'


# ── sheets ────────────────────────────────────────────────────────────────────
def _results_sheet(ws, exams, test, questions, done, rest, multi):
    scores = [x['result'].score or 0 for x in done]
    passed = sum(1 for x in done if x['result'].passed and not x['pending'])
    classes = ', '.join(e.class_room.name for e in exams)
    when = timezone.localtime(exams[0].start_time).strftime('%d.%m.%Y')
    total = done[0]['result'].total_points if done else sum(q.points for q in questions)
    info = [
        f'Sinif: {classes}   ·   Tarix: {when}   ·   Sual: {len(questions)}   ·   '
        f'Maksimum bal: {_num(total)}   ·   Keçid balı: {_num(test.passing_score)}',
        (f'İştirakçı: {len(done)}   ·   Orta bal: {_num(sum(scores) / len(scores))}   ·   '
         f'Ən yüksək: {_num(max(scores))}   ·   Ən aşağı: {_num(min(scores))}   ·   '
         f'Keçdi: {passed} ({round(passed / len(done) * 100)}%)') if done else 'Hələ təhvil verən yoxdur',
    ]
    headers = ['Yer', 'Ad Soyad', 'İstifadəçi adı'] + (['Sinif'] if multi else []) + \
              ['Düz', 'Səhv', 'Boş', 'Bal', 'Faiz', 'Nəticə', 'Təhvil vaxtı', 'Növ']
    row = _title_block(ws, len(headers), exams[0].title or test.title, info)
    head_row = row
    _header(ws, row, headers)
    ws.row_dimensions[row].height = 22
    col = {h: i + 1 for i, h in enumerate(headers)}

    place, last_score = 0, None
    for n, x in enumerate(done + rest, start=1):
        r = x['result']
        row += 1
        submitted = r.status == ExamResult.STATUS_SUBMITTED
        if submitted and r.score != last_score:
            place, last_score = n, r.score
        if x['pending']:
            verdict, v_color, v_bg = 'Yoxlanılır', AMBER, AMBER_BG
        elif not submitted:
            verdict, v_color, v_bg = 'Davam edir', AMBER, AMBER_BG
        elif r.passed:
            verdict, v_color, v_bg = 'Keçdi', GREEN_DARK, GREEN_BG
        else:
            verdict, v_color, v_bg = 'Keçmədi', RED, RED_BG
        values = {
            'Yer': place if submitted else '—', 'Ad Soyad': x['name'], 'İstifadəçi adı': x['username'],
            'Sinif': x['class'], 'Düz': x['correct'], 'Səhv': x['wrong'], 'Boş': x['empty'],
            'Bal': _num(r.score) if submitted else None,
            'Faiz': (r.percentage or 0) / 100 if submitted else None,
            'Nəticə': verdict,
            'Təhvil vaxtı': timezone.localtime(r.submitted_at).strftime('%d.%m.%Y %H:%M') if r.submitted_at else '',
            'Növ': 'Kağız' if r.is_offline else 'Onlayn',
        }
        zebra = _fill(ZEBRA) if n % 2 == 0 else None
        for h, i in col.items():
            c = ws.cell(row=row, column=i, value=values[h])
            c.border = BOX
            c.font = _font()
            c.alignment = LEFT if h in ('Ad Soyad', 'İstifadəçi adı', 'Sinif') else CENTER
            if zebra:
                c.fill = zebra
        ws.cell(row=row, column=col['Düz']).font = _font(GREEN, bold=True)
        ws.cell(row=row, column=col['Səhv']).font = _font(RED, bold=True)
        ws.cell(row=row, column=col['Boş']).font = _font(MUTED)
        ws.cell(row=row, column=col['Bal']).font = _font(INK, bold=True)
        ws.cell(row=row, column=col['Faiz']).number_format = '0%'
        v = ws.cell(row=row, column=col['Nəticə'])
        v.font, v.fill = _font(v_color, bold=True), _fill(v_bg)
        if submitted and place <= 3:
            ws.cell(row=row, column=col['Yer']).font = _font(NAVY, bold=True)

    if row > head_row:
        pct = get_column_letter(col['Faiz'])
        ws.conditional_formatting.add(f'{pct}{head_row + 1}:{pct}{row}', DataBarRule(
            start_type='num', start_value=0, end_type='num', end_value=1, color='60A5FA', showValue=True))
        ws.auto_filter.ref = f'A{head_row}:{get_column_letter(len(headers))}{row}'
    ws.freeze_panes = ws.cell(row=head_row + 1, column=3)
    _widths(ws, [6, 26, 18] + ([10] if multi else []) + [7, 7, 7, 8, 14, 13, 17, 9])


def _answers_sheet(ws, test, questions, choices, done, rest, multi):
    lead = ['Yer', 'Ad Soyad'] + (['Sinif'] if multi else [])
    first_q = len(lead) + 1
    width = len(lead) + len(questions)
    row = _title_block(ws, max(width, 6), 'Şagird cavabları',
                       ['Yaşıl — düz, qırmızı — səhv, boş xana — cavab verilməyib.  '
                        '✓ / ½ / ✗ — əl ilə yoxlanan suallar, ? — hələ yoxlanılmayıb.'])
    _header(ws, row, lead + [i for i in range(1, len(questions) + 1)])
    head_row = row
    ws.row_dimensions[row].height = 20

    # the key, right under the numbers
    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(lead))
    c = ws.cell(row=row, column=1, value='Düzgün cavab')
    c.font, c.fill, c.alignment = _font(NAVY, bold=True), _fill(BLUE_BG), Alignment(horizontal='right', vertical='center', indent=1)
    for j, q in enumerate(questions):
        k = ws.cell(row=row, column=first_q + j, value=_key_text(q, choices[q.id]))
        k.font, k.fill, k.alignment, k.border = _font(NAVY, bold=True), _fill(BLUE_BG), CENTER, BOX
    for i in range(1, len(lead) + 1):
        ws.cell(row=row, column=i).border = BOX

    place, last_score = 0, None
    for n, x in enumerate(done + rest, start=1):
        r = x['result']
        row += 1
        submitted = r.status == ExamResult.STATUS_SUBMITTED
        if submitted and r.score != last_score:
            place, last_score = n, r.score
        lead_values = [place if submitted else '—', x['name']] + ([x['class']] if multi else [])
        for i, value in enumerate(lead_values, start=1):
            c = ws.cell(row=row, column=i, value=value)
            c.border, c.font = BOX, _font()
            c.alignment = CENTER if i == 1 else LEFT
        for j, (text, state) in enumerate(x['cells']):
            c = ws.cell(row=row, column=first_q + j, value=text or None)
            c.border, c.alignment = BOX, CENTER
            if state == 'right':
                c.font, c.fill = _font(GREEN_DARK, bold=True), _fill(GREEN_BG)
            elif state in ('wrong',):
                c.font, c.fill = _font(RED, bold=True), _fill(RED_BG)
            elif state in ('part', 'pending'):
                c.font, c.fill = _font(AMBER, bold=True), _fill(AMBER_BG)
            else:
                c.fill = _fill('F1F5F9')

    ws.freeze_panes = ws.cell(row=head_row + 2, column=first_q)
    _widths(ws, [6, 26] + ([10] if multi else []) + [5] * len(questions))
    for j, q in enumerate(questions):
        if q.question_type not in Question.MCQ_TYPES:   # typed answers need a little more room
            ws.column_dimensions[get_column_letter(first_q + j)].width = 9


def _analysis_sheet(ws, questions, choices, done):
    n_letters = max([len(choices[q.id]) for q in questions if q.question_type in Question.MCQ_TYPES] or [0])
    n_letters = min(n_letters, len(LETTERS))
    headers = ['Sual', 'Növ', 'Düzgün cavab', 'Düz cavab', 'Çətinlik'] + list(LETTERS[:n_letters]) + ['Boş']
    row = _title_block(ws, len(headers), 'Sual təhlili', [
        f'{len(done)} şagirdin cavabı əsasında.  Variant sütunları — hər variantı seçənlərin faizi; '
        'yaşıl — düzgün variant, qırmızı — ən çox seçilən səhv variant.'])
    _header(ws, row, headers)
    head_row = row
    type_names = {'short_answer': 'Qısa cavab', 'numeric': 'Rəqəm', 'open': 'Açıq', 'long_answer': 'Açıq'}
    total = len(done) or 1

    for n, q in enumerate(questions, start=1):
        row += 1
        q_choices = choices[q.id][:n_letters] if q.question_type in Question.MCQ_TYPES else []
        answers = [x['given'].get(q.id) for x in done]
        right = sum(1 for a in answers if a and a.is_correct)
        if q.question_type in Question.MANUAL_GRADE_TYPES:     # partial credit counts as its share
            right_share = sum(a.points_earned / q.points for a in answers if a and q.points) / total
        else:
            right_share = right / total
        blank = sum(1 for x in done if x['cells'][n - 1][1] == 'blank') / total
        level, lv_color = (('Asan', GREEN_DARK) if right_share >= 0.7 else
                           ('Orta', AMBER) if right_share >= 0.4 else ('Çətin', RED))
        values = [n, 'Test' if q.question_type in Question.MCQ_TYPES else type_names.get(q.question_type, ''),
                  _key_text(q, choices[q.id]), right_share, level]
        for i, value in enumerate(values, start=1):
            c = ws.cell(row=row, column=i, value=value)
            c.border, c.alignment, c.font = BOX, CENTER, _font()
        ws.cell(row=row, column=4).number_format = '0%'
        ws.cell(row=row, column=5).font = _font(lv_color, bold=True)
        ws.cell(row=row, column=3).font = _font(NAVY, bold=True)

        shares = [sum(1 for a in answers if a and a.selected_answer_id == ch.id) / total for ch in q_choices]
        wrong_shares = [s for s, ch in zip(shares, q_choices) if not ch.is_correct]
        top_wrong = max(wrong_shares) if wrong_shares else 0
        for i in range(n_letters):
            c = ws.cell(row=row, column=6 + i)
            c.border, c.alignment = BOX, CENTER
            if i >= len(q_choices):
                c.value, c.font = '—', _font(LINE)
                continue
            share, ch = shares[i], q_choices[i]
            c.value, c.number_format = share, '0%'
            if ch.is_correct:
                c.fill, c.font = _fill(GREEN_BG), _font(GREEN_DARK, bold=True)
            elif share and share == top_wrong and share >= 0.15:
                c.fill, c.font = _fill(RED_BG), _font(RED, bold=True)
            else:
                c.font = _font(MUTED if not share else INK)
        b = ws.cell(row=row, column=6 + n_letters, value=blank)
        b.border, b.alignment, b.number_format, b.font = BOX, CENTER, '0%', _font(MUTED)
        if n % 2 == 0:
            for i in (1, 2):
                ws.cell(row=row, column=i).fill = _fill(ZEBRA)

    if row > head_row:
        ws.conditional_formatting.add(f'D{head_row + 1}:D{row}', DataBarRule(
            start_type='num', start_value=0, end_type='num', end_value=1, color='4ADE80', showValue=True))
    ws.freeze_panes = ws.cell(row=head_row + 1, column=2)
    _widths(ws, [7, 12, 14, 14, 11] + [8] * n_letters + [8])


def build_results_workbook(exams):
    """exams: the same test given to one or more classes (same teacher). Returns .xlsx bytes."""
    test, questions, choices, done, rest = _collect(exams)
    multi = len(exams) > 1
    wb = Workbook()
    ws1 = wb.active
    ws1.title = 'Nəticələr'
    ws1.sheet_properties.tabColor = '16A34A'
    _results_sheet(ws1, exams, test, questions, done, rest, multi)

    ws2 = wb.create_sheet('Cavablar')
    ws2.sheet_properties.tabColor = 'DC2626'
    _answers_sheet(ws2, test, questions, choices, done, rest, multi)

    ws3 = wb.create_sheet('Sual təhlili')
    ws3.sheet_properties.tabColor = '2563EB'
    _analysis_sheet(ws3, questions, choices, done)

    for ws in (ws1, ws2, ws3):
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
